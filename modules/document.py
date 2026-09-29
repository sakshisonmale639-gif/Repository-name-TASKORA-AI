from docx import Document
import os


def create_docx(title, content):

    os.makedirs("generated/documents", exist_ok=True)

    filename = f"{title.replace(' ', '_')}.docx"
    path = os.path.join("generated/documents", filename)

    document = Document()

    document.add_heading(title, 0)

    for paragraph in content.split("\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph)

    document.save(path)

    return path