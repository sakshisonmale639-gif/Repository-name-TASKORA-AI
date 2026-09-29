from pptx import Presentation
import os


def create_presentation(title, content, slide_count=5):

    os.makedirs("generated/presentations", exist_ok=True)

    filename = f"{title.replace(' ', '_')}.pptx"
    path = os.path.join("generated/presentations", filename)

    presentation = Presentation()

    # Title slide
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[0]
    )

    slide.shapes.title.text = title

    if slide.placeholders:
        slide.placeholders[1].text = "Created with TASKORA AI"

    paragraphs = [
        p.strip()
        for p in content.split("\n")
        if p.strip()
    ]

    for i in range(slide_count - 1):

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[1]
        )

        slide.shapes.title.text = f"{title} - Slide {i + 2}"

        text_frame = slide.placeholders[1].text_frame

        if paragraphs:
            text_frame.text = paragraphs[
                i % len(paragraphs)
            ]
        else:
            text_frame.text = "TASKORA AI"

    presentation.save(path)

    return path