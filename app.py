import os
import re
from pathlib import Path

import streamlit as st

# ============================================================
# GEMINI API KEY
# ============================================================

try:
    if "GEMINI_API_KEY" in st.secrets:
        os.environ["GEMINI_API_KEY"] = st.secrets["GEMINI_API_KEY"]
except Exception:
    pass

# ---------------------------------------------------------
# TASKORA - MAIN APPLICATION
# ---------------------------------------------------------

st.set_page_config(
    page_title="TASKORA AI",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---------------------------------------------------------
# PATHS
# ---------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent

GENERATED_DIR = BASE_DIR / "generated"
IMAGE_DIR = GENERATED_DIR / "images"
DOCUMENT_DIR = GENERATED_DIR / "documents"
PRESENTATION_DIR = GENERATED_DIR / "presentations"
PDF_DIR = GENERATED_DIR / "pdfs"

for folder in [
    IMAGE_DIR,
    DOCUMENT_DIR,
    PRESENTATION_DIR,
    PDF_DIR,
]:
    folder.mkdir(parents=True, exist_ok=True)


# ---------------------------------------------------------
# OPTIONAL TASKORA MODULES
# ---------------------------------------------------------

try:
    from modules.image_generator import generate_image
    IMAGE_GENERATOR_AVAILABLE = True
except Exception as e:
    generate_image = None
    IMAGE_GENERATOR_AVAILABLE = False
    IMAGE_GENERATOR_ERROR = str(e)


try:
    from agent import run_agent
    AGENT_AVAILABLE = True
except Exception:
    run_agent = None
    AGENT_AVAILABLE = False


# ---------------------------------------------------------
# SESSION STATE
# ---------------------------------------------------------

if "page" not in st.session_state:
    st.session_state.page = "🎨 Image Creator"

if "tasks" not in st.session_state:
    st.session_state.tasks = []

if "generated_image" not in st.session_state:
    st.session_state.generated_image = None

if "assistant_messages" not in st.session_state:
    st.session_state.assistant_messages = []


# ---------------------------------------------------------
# CUSTOM CSS
# ---------------------------------------------------------

st.markdown(
    """
<style>

.stApp {
    background:
        radial-gradient(circle at 10% 10%, rgba(123, 63, 255, 0.25), transparent 30%),
        radial-gradient(circle at 90% 20%, rgba(0, 180, 255, 0.20), transparent 30%),
        radial-gradient(circle at 50% 100%, rgba(255, 0, 180, 0.18), transparent 35%),
        linear-gradient(135deg, #080b22 0%, #111640 45%, #21103d 100%);
    color: white;
}

/* Main content */
.block-container {
    max-width: 1200px;
    padding-top: 2rem;
    padding-bottom: 4rem;
}

/* Sidebar */
section[data-testid="stSidebar"] {
    background:
        linear-gradient(
            180deg,
            #080b22 0%,
            #10143a 45%,
            #17113b 100%
        );
    border-right: 1px solid rgba(255,255,255,0.08);
}

section[data-testid="stSidebar"] * {
    color: #f5f5ff !important;
}

/* Headings */
h1 {
    font-size: 3rem !important;
    font-weight: 800 !important;
    background: linear-gradient(
        90deg,
        #ffffff,
        #a78bfa,
        #60a5fa
    );
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
}

h2, h3 {
    color: white !important;
}

/* Cards */
.taskora-card {
    padding: 28px;
    border-radius: 24px;
    background: rgba(255,255,255,0.075);
    border: 1px solid rgba(255,255,255,0.12);
    box-shadow: 0 20px 50px rgba(0,0,0,0.25);
    backdrop-filter: blur(16px);
    margin-bottom: 24px;
}

.hero-card {
    padding: 45px;
    border-radius: 30px;
    background:
        linear-gradient(
            135deg,
            rgba(92, 45, 170, 0.65),
            rgba(29, 78, 216, 0.45)
        );
    border: 1px solid rgba(255,255,255,0.14);
    box-shadow: 0 25px 70px rgba(0,0,0,0.35);
    margin-bottom: 30px;
}

.hero-title {
    font-size: 42px;
    font-weight: 850;
    color: white;
    margin-bottom: 10px;
}

.hero-text {
    color: #e0e7ff;
    font-size: 18px;
}

.small-label {
    color: #a5b4fc;
    font-weight: 700;
    letter-spacing: 1px;
    text-transform: uppercase;
}

.success-box {
    padding: 16px;
    border-radius: 14px;
    background: rgba(34,197,94,0.15);
    border: 1px solid rgba(34,197,94,0.3);
}

.error-box {
    padding: 16px;
    border-radius: 14px;
    background: rgba(239,68,68,0.15);
    border: 1px solid rgba(239,68,68,0.3);
}

/* Buttons */
.stButton > button {
    border-radius: 14px;
    border: none;
    padding: 0.65rem 1.2rem;
    font-weight: 700;
    background: linear-gradient(
        90deg,
        #7c3aed,
        #2563eb
    );
    color: white;
    box-shadow: 0 8px 25px rgba(76,29,149,0.35);
}

.stButton > button:hover {
    transform: translateY(-2px);
    box-shadow: 0 12px 30px rgba(76,29,149,0.5);
}

/* Inputs */
.stTextInput input,
.stTextArea textarea {
    border-radius: 14px !important;
    background: rgba(255,255,255,0.95) !important;
    color: #111827 !important;
}

/* Radio */
div[role="radiogroup"] label {
    padding: 8px 5px;
}

/* Divider */
hr {
    border-color: rgba(255,255,255,0.1);
}

</style>
""",
    unsafe_allow_html=True,
)


# ---------------------------------------------------------
# SIDEBAR
# ---------------------------------------------------------

with st.sidebar:

    st.markdown(
        """
        <div style="text-align:center; padding:20px 0 25px 0;">
            <div style="font-size:58px;">🧠</div>
            <div style="
                font-size:30px;
                font-weight:850;
                color:#8ab4ff;
            ">
                TASKORA
            </div>
            <div style="color:#a5b4fc;">
                AI Productivity Studio
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.divider()

    st.markdown(
        '<div class="small-label">Workspace</div>',
        unsafe_allow_html=True,
    )

    pages = [
        "🏠 Dashboard",
        "🤖 AI Assistant",
        "✅ Tasks",
        "📄 Documents",
        "📊 Presentations",
        "📕 PDFs",
        "🎨 Image Creator",
    ]

    selected_page = st.radio(
        "Navigate",
        pages,
        index=pages.index(st.session_state.page),
        label_visibility="collapsed",
    )

    st.session_state.page = selected_page

    st.divider()

    st.markdown(
        """
        <div style="
            text-align:center;
            color:#94a3b8;
            padding:20px 0;
        ">
            <b>TASKORA AI</b><br><br>
            Create • Think • Plan • Design • Automate
        </div>
        """,
        unsafe_allow_html=True,
    )


# ---------------------------------------------------------
# HELPER FUNCTIONS
# ---------------------------------------------------------

def safe_filename(text):
    text = re.sub(r"[^a-zA-Z0-9_-]+", "_", text)
    return text[:80] or "taskora_file"


def create_docx(title, content):
    from docx import Document

    filename = safe_filename(title) + ".docx"
    path = DOCUMENT_DIR / filename

    document = Document()

    document.add_heading(title, 0)

    for paragraph in content.split("\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())

    document.save(path)

    return path


def create_pdf(title, content):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
    )

    filename = safe_filename(title) + ".pdf"
    path = PDF_DIR / filename

    styles = getSampleStyleSheet()

    pdf = SimpleDocTemplate(
        str(path),
        pagesize=letter,
    )

    story = []

    story.append(
        Paragraph(
            title,
            styles["Title"],
        )
    )

    story.append(Spacer(1, 20))

    for paragraph in content.split("\n"):
        if paragraph.strip():
            story.append(
                Paragraph(
                    paragraph.strip(),
                    styles["BodyText"],
                )
            )
            story.append(Spacer(1, 10))

    pdf.build(story)

    return path


def create_presentation(title, content):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()

    # Title slide
    slide = prs.slides.add_slide(
        prs.slide_layouts[0]
    )

    slide.shapes.title.text = title

    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = (
            "Created with TASKORA AI"
        )

    # Content slides
    paragraphs = [
        p.strip()
        for p in content.split("\n")
        if p.strip()
    ]

    # Create a slide approximately every 3 paragraphs
    chunks = [
        paragraphs[i:i + 3]
        for i in range(0, len(paragraphs), 3)
    ]

    if not chunks:
        chunks = [["Content generated by TASKORA AI"]]

    for index, chunk in enumerate(chunks, start=1):

        slide = prs.slides.add_slide(
            prs.slide_layouts[1]
        )

        slide.shapes.title.text = (
            f"{title} — Part {index}"
        )

        body = slide.placeholders[1].text_frame

        body.clear()

        for item in chunk:
            paragraph = body.add_paragraph()
            paragraph.text = item
            paragraph.font.size = Pt(22)

    filename = safe_filename(title) + ".pptx"
    path = PRESENTATION_DIR / filename

    prs.save(path)

    return path


def ai_generate(prompt):
    if AGENT_AVAILABLE:
        try:
            return run_agent(prompt)
        except Exception as e:
            return (
                "TASKORA could not reach the AI service.\n\n"
                f"Error: {e}"
            )

    return (
        "TASKORA AI Assistant\n\n"
        "The AI agent module is not currently available.\n"
        "Please check agent.py."
    )


# ---------------------------------------------------------
# DASHBOARD
# ---------------------------------------------------------

if st.session_state.page == "🏠 Dashboard":

    st.markdown(
        """
        <div class="hero-card">
            <div class="hero-title">
                Welcome to TASKORA ✨
            </div>
            <div class="hero-text">
                Your AI productivity studio for creating,
                planning, designing and automating work.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    col1, col2, col3, col4 = st.columns(4)

    with col1:
        st.metric("Tasks", len(st.session_state.tasks))

    with col2:
        st.metric("Documents", len(list(DOCUMENT_DIR.glob("*"))))

    with col3:
        st.metric(
            "Presentations",
            len(list(PRESENTATION_DIR.glob("*"))),
        )

    with col4:
        st.metric(
            "Images",
            len(list(IMAGE_DIR.glob("*"))),
        )

    st.markdown(
        """
        <div class="taskora-card">
            <h2>🚀 What can TASKORA do?</h2>
            <p>
            🤖 AI Assistant &nbsp; • &nbsp;
            📄 Documents &nbsp; • &nbsp;
            📊 Presentations &nbsp; • &nbsp;
            📕 PDFs &nbsp; • &nbsp;
            🎨 AI Images &nbsp; • &nbsp;
            ✅ Tasks
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )


# ---------------------------------------------------------
# AI ASSISTANT
# ---------------------------------------------------------

elif st.session_state.page == "🤖 AI Assistant":

    st.markdown(
        """
        <div class="hero-card">
            <div class="hero-title">
                🤖 TASKORA AI Assistant
            </div>
            <div class="hero-text">
                Ask TASKORA anything and let your AI
                productivity assistant help you.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    prompt = st.text_area(
        "What do you want TASKORA to do?",
        placeholder=(
            "Example: Explain artificial intelligence "
            "in simple terms..."
        ),
        height=160,
    )

    if st.button("✨ Ask TASKORA"):

        if not prompt.strip():
            st.warning("Please enter a request.")

        else:
            with st.spinner("TASKORA is thinking..."):

                answer = ai_generate(prompt)

            st.markdown(
                '<div class="taskora-card">',
                unsafe_allow_html=True,
            )

            st.markdown("### 💡 TASKORA")

            st.write(answer)

            st.markdown("</div>", unsafe_allow_html=True)


# ---------------------------------------------------------
# TASKS
# ---------------------------------------------------------

elif st.session_state.page == "✅ Tasks":

    st.markdown(
        """
        <div class="hero-card">
            <div class="hero-title">
                ✅ Task Manager
            </div>
            <div class="hero-text">
                Organize your work and keep track of your goals.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    task = st.text_input(
        "New task",
        placeholder="Example: Finish TASKORA PPT module",
    )

    if st.button("➕ Add Task"):

        if task.strip():
            st.session_state.tasks.append(
                {
                    "task": task.strip(),
                    "done": False,
                }
            )

            st.success("Task added!")

    st.markdown("### Your Tasks")

    if not st.session_state.tasks:

        st.info("No tasks yet.")

    else:

        for index, item in enumerate(
            st.session_state.tasks
        ):

            col1, col2 = st.columns([5, 1])

            with col1:

                checked = st.checkbox(
                    item["task"],
                    value=item["done"],
                    key=f"task_{index}",
                )

                st.session_state.tasks[index]["done"] = checked

            with col2:

                if st.button(
                    "🗑️",
                    key=f"delete_{index}",
                ):

                    st.session_state.tasks.pop(index)

                    st.rerun()


# ---------------------------------------------------------
# DOCUMENTS
# ---------------------------------------------------------

elif st.session_state.page == "📄 Documents":

    st.markdown(
        """
        <div class="hero-card">
            <div class="hero-title">
                📄 Document Creator
            </div>
            <div class="hero-text">
                Create Microsoft Word documents instantly.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    title = st.text_input(
        "Document title",
        value="TASKORA Document",
    )

    content = st.text_area(
        "Document content",
        height=300,
        placeholder="Write your document content here...",
    )

    if st.button("📄 Create DOCX"):

        if not content.strip():
            st.warning("Please enter some content.")

        else:

            try:

                path = create_docx(
                    title,
                    content,
                )

                st.success(
                    "Document created successfully!"
                )

                with open(path, "rb") as file:

                    st.download_button(
                        "⬇️ Download DOCX",
                        file,
                        file_name=path.name,
                        mime=(
                            "application/vnd.openxmlformats-"
                            "officedocument.wordprocessingml.document"
                        ),
                    )

            except Exception as e:

                st.error(
                    f"Document creation failed: {e}"
                )


# ---------------------------------------------------------
# PRESENTATIONS
# ---------------------------------------------------------

elif st.session_state.page == "📊 Presentations":

    st.markdown(
        """
        <div class="hero-card">
            <div class="hero-title">
                📊 TASKORA Presentation Creator
            </div>
            <div class="hero-text">
                Create PowerPoint presentations from a simple idea.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    title = st.text_input(
        "Presentation title",
        value="AI in Education",
    )

    content = st.text_area(
        "Presentation content",
        height=300,
        placeholder=(
            "Example:\n"
            "Artificial intelligence is changing education.\n"
            "Personalized learning helps students.\n"
            "AI tutors can provide instant feedback.\n"
            "Teachers can automate repetitive tasks.\n"
            "The future of education will be more personalized."
        ),
    )

    if st.button("📊 Create PowerPoint"):

        if not content.strip():
            st.warning("Please enter presentation content.")

        else:

            try:

                path = create_presentation(
                    title,
                    content,
                )

                st.success(
                    "PowerPoint created successfully!"
                )

                with open(path, "rb") as file:

                    st.download_button(
                        "⬇️ Download PPTX",
                        file,
                        file_name=path.name,
                        mime=(
                            "application/vnd.openxmlformats-"
                            "officedocument.presentationml.presentation"
                        ),
                    )

            except Exception as e:

                st.error(
                    f"PowerPoint creation failed: {e}"
                )


# ---------------------------------------------------------
# PDF
# ---------------------------------------------------------

elif st.session_state.page == "📕 PDFs":

    st.markdown(
        """
        <div class="hero-card">
            <div class="hero-title">
                📕 PDF Creator
            </div>
            <div class="hero-text">
                Turn your content into downloadable PDF files.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    title = st.text_input(
        "PDF title",
        value="TASKORA PDF",
    )

    content = st.text_area(
        "PDF content",
        height=300,
        placeholder="Enter your PDF content...",
    )

    if st.button("📕 Create PDF"):

        if not content.strip():
            st.warning("Please enter some content.")

        else:

            try:

                path = create_pdf(
                    title,
                    content,
                )

                st.success(
                    "PDF created successfully!"
                )

                with open(path, "rb") as file:

                    st.download_button(
                        "⬇️ Download PDF",
                        file,
                        file_name=path.name,
                        mime="application/pdf",
                    )

            except Exception as e:

                st.error(
                    f"PDF creation failed: {e}"
                )


# ---------------------------------------------------------
# IMAGE CREATOR
# ---------------------------------------------------------

# ============================================================
# IMAGE CREATOR
# ============================================================

if selected_page == "Image Creator":

    st.markdown(
        """
        <div class="hero">
            <h1>🎨 Bring Your Ideas to Life ✨</h1>
            <p>
                Describe an image and TASKORA will generate it
                using Gemini AI.
            </p>
        </div>
        """,
        unsafe_allow_html=True
    )

    st.subheader("✨ Describe your image")

    image_prompt = st.text_area(
        "Image description",
        placeholder=(
            "Example: A futuristic AI assistant helping "
            "a college student organize tasks in a modern "
            "digital workspace, purple and blue neon "
            "lighting, cinematic, highly detailed..."
        ),
        height=150,
        key="image_prompt"
    )

    # --------------------------------------------------------
    # GENERATE IMAGE BUTTON
    # --------------------------------------------------------

    if st.button(
        "✨ Generate Image",
        type="primary",
        use_container_width=False
    ):

        if not image_prompt.strip():

            st.warning(
                "⚠️ Please describe the image you want to generate."
            )

        else:

            with st.spinner(
                "🎨 TASKORA is creating your image..."
            ):

                try:

                   
                    image_path = generate_image(
                        image_prompt.strip()
                    )

                    st.success(
                        "✅ Image generated successfully!"
                    )

                    # Display generated image
                    st.image(
                        image_path,
                        caption="Generated by TASKORA AI",
                        use_container_width=True
                    )

                    # Download button
                    with open(image_path, "rb") as image_file:

                        st.download_button(
                            label="⬇️ Download Image",
                            data=image_file.read(),
                            file_name="taskora_generated_image.png",
                            mime="image/png"
                        )

                except Exception as error:

                    st.error(
                        "❌ Image generation failed."
                    )

                    st.exception(error)


# ---------------------------------------------------------
# FOOTER
# ---------------------------------------------------------

st.markdown(
    """
    <div style="
        text-align:center;
        padding:45px 0 10px 0;
        color:#94a3b8;
    ">
        🧠 <b>TASKORA AI</b> • AI Productivity Studio
        <br><br>
        Create • Think • Plan • Design • Automate
    </div>
    """,
    unsafe_allow_html=True,
)